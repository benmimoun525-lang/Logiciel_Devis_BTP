import os
import io
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Body
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
import google.generativeai as genai

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

current_key_index = 0

def get_api_keys_pool():
    keys = []
    for key, value in os.environ.items():
        if key.startswith("GEMINI_KEY_") or key == "GEMINI_API_KEY":
            if value and value.strip():
                keys.append(value.strip())
    return keys

MODELS_PRIORITY = ['gemini-1.5-flash']

@app.get("/")
def read_root():
    keys_count = len(get_api_keys_pool())
    return {"status": "ok", "message": f"Serveur BTP opérationnel ({keys_count} clés). Mode : Extraction PU HT + Formules Excel."}

@app.post("/chiffrer-page")
async def chiffrer_page(
    file: UploadFile = File(None),
    texte_descriptif: str = Form(None),
    page_num: int = Form(1),
    start_index: int = Form(1)
):
    global current_key_index
    try:
        keys_pool = get_api_keys_pool()
        if not keys_pool:
            raise HTTPException(status_code=500, detail="Aucune clé API configurée.")

        prompt_base = (
            f"Tu es un expert métreur BTP en Algérie.\n"
            f"Analyse CETTE PAGE (Page {page_num}) et extrait tous les articles.\n"
            f"Commence la numérotation à {start_index}.\n\n"
            "FORMAT DE RÉPONSE STRICT :\n"
            "[\n"
            f"  {{\"n\": {start_index}, \"d\": \"Désignation précise\", \"u\": \"m3\", \"q\": 10, \"pu\": 12000}}\n"
            "]\n\n"
            "CONSIGNES :\n"
            "- 'q' = Quantité (mets 1 si vide).\n"
            "- 'pu' = Estime un Prix Unitaire HT réaliste en DZD.\n"
            "- NE CALCULE AUCUN TOTAL.\n"
            "- RÈGLE ABSOLUE : Renvoie UNIQUEMENT le tableau JSON, aucun texte avant, aucun texte après."
        )

        contents_list = [prompt_base]
        if texte_descriptif and texte_descriptif.strip():
            contents_list.append(f"\n--- DESCRIPTIF ---\n{texte_descriptif.strip()}")

        if file:
            file_bytes = await file.read()
            if file_bytes:
                content_type = file.content_type or "image/jpeg"
                contents_list.append({"mime_type": content_type, "data": file_bytes})

        total_keys = len(keys_pool)
        last_error = ""

        for attempt in range(total_keys):
            selected_key_idx = (current_key_index + attempt) % total_keys
            genai.configure(api_key=keys_pool[selected_key_idx])
            key_failed_due_to_quota = False

            for model_name in MODELS_PRIORITY:
                try:
                    model = genai.GenerativeModel(model_name=model_name)
                    
                    # Configuration passée sous forme de dictionnaire simple (ultra compatible)
                    response = model.generate_content(
                        contents_list,
                        generation_config={"max_output_tokens": 2048, "temperature": 0.0}
                    )

                    raw_text = response.text or "[]"
                    clean_json = raw_text.replace("```json", "").replace("```", "").strip()

                    current_key_index = (selected_key_idx + 1) % total_keys
                    time.sleep(1.5)

                    return JSONResponse(content={"page": page_num, "raw_json": clean_json})

                except Exception as inner_e:
                    err_msg = str(inner_e).lower()
                    last_error = err_msg
                    if "429" in err_msg or "quota" in err_msg:
                        key_failed_due_to_quota = True
                        break
                    else:
                        continue

            if key_failed_due_to_quota:
                time.sleep(1)
                continue

        raise HTTPException(status_code=429, detail=f"Détail de l'erreur : {last_error}")

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/exporter-excel")
async def exporter_excel(payload: dict = Body(...)):
    try:
        nom_client = payload.get("nom_client", "Client")
        articles = payload.get("articles", [])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Devis Estimatif"

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        ws['A1'] = f"DEVIS QUANTITATIF ET ESTIMATIF : {nom_client}"
        ws['A1'].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        
        headers = ["N°", "Désignation des Travaux", "Unité", "Quantité", "P.U HT (DZD)", "Montant HT (DZD)"]
        ws.append([]) 
        ws.append(headers) 
        
        for col_num in range(1, 7):
            cell = ws.cell(row=3, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        start_row = 4

        for idx, art in enumerate(articles):
            row_idx = start_row + idx
            num = art.get("n", idx + 1)
            des = art.get("d", "Article sans désignation")
            uni = art.get("u", "U")
            qte = float(art.get("q", 1) or 1)
            pu = float(art.get("pu", 0) or 0)

            ws.cell(row=row_idx, column=1, value=num).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=des)
            ws.cell(row=row_idx, column=3, value=uni).alignment = Alignment(horizontal="center")
            
            ws.cell(row=row_idx, column=4, value=qte).number_format = '#,##0.00'
            ws.cell(row=row_idx, column=5, value=pu).number_format = '#,##0.00'
            
            ws.cell(row=row_idx, column=6, value=f"=D{row_idx}*E{row_idx}").number_format = '#,##0.00'

            for c in range(1, 7):
                ws.cell(row=row_idx, column=c).border = thin_border

        last_data_row = start_row + len(articles) - 1
        if len(articles) == 0:
            last_data_row = start_row

        total_row = last_data_row + 2

        ws.cell(row=total_row, column=5, value="TOTAL HT :").font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=f"=SUM(F{start_row}:F{last_data_row})").font = Font(bold=True)
        ws.cell(row=total_row, column=6).number_format = '#,##0.00 DZD'

        ws.cell(row=total_row+1, column=5, value="TVA (19%) :")
        ws.cell(row=total_row+1, column=6, value=f"=F{total_row}*0.19")
        ws.cell(row=total_row+1, column=6).number_format = '#,##0.00 DZD'

        ws.cell(row=total_row+2, column=5, value="TOTAL TTC :").font = Font(bold=True, color="1E3A8A")
        ws.cell(row=total_row+2, column=6, value=f"=F{total_row}+F{total_row+1}").font = Font(bold=True, color="1E3A8A")
        ws.cell(row=total_row+2, column=6).number_format = '#,##0.00 DZD'

        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 60
        ws.column_dimensions['C'].width = 8
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 22

        output_stream = io.BytesIO()
        wb.save(output_stream)
        output_stream.seek(0)

        nom_fichier_propre = nom_client.replace(' ', '_').replace('/', '_')
        return StreamingResponse(
            output_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={'Content-Disposition': f'attachment; filename="Devis_Estime_{nom_fichier_propre}.xlsx"'}
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))